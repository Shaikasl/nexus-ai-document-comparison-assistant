
import os
import re
import time
from io import BytesIO

import streamlit as st
from dotenv import load_dotenv
from pypdf import PdfReader
from google import genai
from docx import Document

load_dotenv()

MODEL_NAME = "gemini-3.8-flash"
MAX_CHARS = 60000

st.set_page_config(
    page_title="Nexus AI | Document Comparison",
    page_icon="📄",
    layout="wide",
    initial_sidebar_state="expanded",
)

# ---------- SESSION STATE ----------

for key, default in {
    "comparison_result": None,
    "original_name": "",
    "revised_name": "",
}.items():
    if key not in st.session_state:
        st.session_state[key] = default


# ---------- UI DESIGN ----------

st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800&display=swap');

:root {
    --bg: #090B18;
    --panel: #14172A;
    --border: #2B2F4B;
    --purple: #A78BFA;
    --cyan: #67E8F9;
    --text: #F5F3FF;
    --muted: #A5AAC6;
}

html, body, [data-testid="stAppViewContainer"] {
    background: var(--bg);
    color: var(--text);
    font-family: 'Inter', sans-serif;
}

[data-testid="stHeader"] {
    background: rgba(9, 11, 24, 0.9);
}

[data-testid="stSidebar"] {
    background: linear-gradient(180deg, #121329, #0B1020);
    border-right: 1px solid var(--border);
}

[data-testid="stSidebar"] * {
    color: var(--text);
}

.block-container {
    max-width: 1450px;
    padding-top: 2rem;
    padding-bottom: 3rem;
}

.brand {
    font-size: 1.35rem;
    font-weight: 800;
    letter-spacing: 1px;
    background: linear-gradient(90deg, #C4B5FD, #67E8F9);
    -webkit-background-clip: text;
    -webkit-text-fill-color: transparent;
}

.eyebrow {
    color: var(--cyan);
    font-size: 0.75rem;
    font-weight: 700;
    letter-spacing: 2px;
    text-transform: uppercase;
}

.hero-title {
    font-size: clamp(2rem, 4vw, 3.4rem);
    font-weight: 800;
    letter-spacing: -1.8px;
    line-height: 1.12;
    margin: 10px 0;
}

.gradient-text {
    background: linear-gradient(90deg, #C4B5FD, #818CF8, #67E8F9);
    -webkit-background-clip: text;
    -webkit-text-fill-color: transparent;
}

.subtext {
    color: var(--muted);
    font-size: 1rem;
    line-height: 1.8;
}

.panel {
    background: linear-gradient(145deg, #181B32, #111426);
    border: 1px solid var(--border);
    border-radius: 18px;
    padding: 22px;
    margin-bottom: 16px;
}

.panel h3 {
    color: var(--text);
    font-size: 1.05rem;
}

div.stButton > button {
    width: 100%;
    border: 1px solid #8172E9;
    border-radius: 10px;
    min-height: 48px;
    font-weight: 700;
    color: white;
    background: linear-gradient(100deg, #7656E8, #5269E8, #277FBB);
}

[data-testid="stFileUploader"] {
    background: #111529;
    border: 1px dashed #6558A8;
    border-radius: 13px;
    padding: 8px;
}

[data-testid="stFileUploader"] section {
    background: transparent;
}

.stTabs [data-baseweb="tab-list"] {
    gap: 10px;
}

.stTabs [data-baseweb="tab"] {
    color: #C4C8E0;
}

hr {
    border-color: var(--border);
}

footer {
    visibility: hidden;
}

@media (max-width: 768px) {
    .hero-title {
        letter-spacing: -0.8px;
    }
    .panel {
        padding: 15px;
    }
}
</style>
""", unsafe_allow_html=True)


# ---------- PDF EXTRACTION ----------

def extract_pdf_text(uploaded_file):
    """Return text, page count, character count and error."""

    try:
        reader = PdfReader(BytesIO(uploaded_file.getvalue()))

        if reader.is_encrypted:
            return None, 0, 0, (
                f"{uploaded_file.name} is password-protected."
            )

        page_count = len(reader.pages)
        page_texts = []

        for page_number, page in enumerate(reader.pages, start=1):
            page_text = page.extract_text() or ""
            if page_text.strip():
                page_texts.append(
                    f"[Page {page_number}]\n{page_text}"
                )

        text = "\n\n".join(page_texts).strip()
        character_count = len(text)

        if not text:
            return None, page_count, 0, (
                f"No readable text found in {uploaded_file.name}. "
                "It may be scanned or empty. OCR may be required."
            )

        return text, page_count, character_count, None

    except Exception:
        return None, 0, 0, (
            f"Unable to read {uploaded_file.name}. "
            "Please upload a valid PDF."
        )


# ---------- AI RESPONSE VALIDATION ----------

REQUIRED_SECTIONS = [
    "DOCUMENT A",
    "DOCUMENT B",
    "SIMILARITIES",
    "DIFFERENCES",
    "KEY DETAILS COMPARISON",
    "OVERALL CONCLUSION",
    "LIMITATIONS",
]


def validate_ai_response(result):
    if not result or not result.strip():
        return False, "The AI returned an empty response."

    missing = [
        section for section in REQUIRED_SECTIONS
        if section.lower() not in result.lower()
    ]

    if missing:
        return False, (
            "The AI response is incomplete. Missing sections: "
            + ", ".join(missing)
        )

    return True, None


# ---------- DOCUMENT COMPARISON ----------


def compare_documents(original_text, revised_text):
    api_key = os.getenv("GEMINI_API_KEY")

    if not api_key:
        return None, "GEMINI_API_KEY is missing. Check your .env file."

    if not original_text or not revised_text:
        return None, "Please upload two readable PDF documents."

    if len(original_text) + len(revised_text) > MAX_CHARS:
        return None, (
            "The documents are too long. Please use shorter PDFs."
        )

    prompt = f"""
You are Nexus AI, an AI Document Comparison Assistant.

Compare both PDF documents using only their extracted text.
Never invent information. Include page numbers when available.

Return exactly these seven sections:

1. DOCUMENT A
Summarize Document A.

2. DOCUMENT B
Summarize Document B.

3. SIMILARITIES
List information present in both documents.

4. DIFFERENCES
List supported additions, removals, changes, and conflicts.

5. KEY DETAILS COMPARISON
Compare dates, names, amounts, deadlines, and requirements.
Use this format:
Field | Document A | Document B | Change

6. OVERALL CONCLUSION
Explain the main changes neutrally and briefly.

7. LIMITATIONS
Mention unreadable pages, missing information, or processing issues.

DOCUMENT A TEXT:
{original_text}

DOCUMENT B TEXT:
{revised_text}
"""

    try:
        client = genai.Client(api_key=api_key)
    except Exception as e:
        return None, f"Could not initialize Gemini: {e}"

    models_to_try = [
        MODEL_NAME,
        "gemini-flash-latest",
    ]

    
    last_error = None
    retryable_status_codes = {429, 500, 502, 503, 504}

    for model_index, model_name in enumerate(models_to_try):
        for attempt in range(2):
            try:
                response = client.models.generate_content(
                    model=model_name,
                    contents=prompt,
                )

                result = (response.text or "").strip()

                if not result:
                    return None, (
                        "Gemini returned an empty response. "
                        "Please try again."
                    )

                valid, error = validate_ai_response(result)

                if not valid:
                    return None, (
                        f"Response validation failed: {error}"
                    )

                return result, None

            except Exception as e:
                last_error = e

                status_code = (
                    getattr(e, "code", None)
                    or getattr(e, "status_code", None)
                )

                error_text = str(e).lower()

                # Retry temporary service or rate-limit errors.
                is_retryable = (
                    status_code in retryable_status_codes
                    or "timeout" in error_text
                    or "timed out" in error_text
                    or "connection reset" in error_text
                )

                if is_retryable and attempt == 0:
                    time.sleep(2)
                    continue

                # If the model name is unavailable, try the fallback.
                model_unavailable = (
                    status_code == 404
                    or "model not found" in error_text
                    or "not found for api version" in error_text
                )

                if model_unavailable:
                    break

                if is_retryable:
                    break

                return None, (
                    f"Gemini API error ({type(e).__name__}): {e}"
                )

            else:
                break

    return None, (
        "Gemini could not complete the comparison after retrying. "
        f"Please try again later. Details: {last_error}"
    )

# ---------- WORD REPORT ----------


def create_word_report(original_name, revised_name, result):
    doc = Document()

    # Report title
    doc.add_heading("Nexus AI", level=0)
    doc.add_heading("AI Document Comparison Report", level=1)

    doc.add_paragraph(f"Document A: {original_name}")
    doc.add_paragraph(f"Document B: {revised_name}")
    doc.add_paragraph(
        "This report compares text extracted from the two uploaded PDFs. "
        "Please verify important details against the original documents."
    )

    # Split the AI response into the seven sections.
    sections = split_comparison_sections(result)

    section_order = [
        "Document A",
        "Document B",
        "Similarities",
        "Differences",
        "Key Details Comparison",
        "Overall Conclusion",
        "Limitations",
    ]

    for section_name in section_order:
        doc.add_heading(section_name, level=1)

        content = sections.get(section_name, "").strip()

        if not content:
            doc.add_paragraph(
                "No content was provided for this section."
            )
            continue

        for line in content.splitlines():
            line = line.strip()

            if not line:
                continue

            # Remove common Markdown formatting.
            clean_line = re.sub(r"^#{1,6}\s*", "", line)
            clean_line = clean_line.strip().strip("*_").strip()

            if clean_line.startswith(("- ", "* ")):
                doc.add_paragraph(
                    clean_line[2:].strip(),
                    style="List Bullet",
                )
            elif re.match(r"^\d+[.)]\s+", clean_line):
                item = re.sub(
                    r"^\d+[.)]\s+",
                    "",
                    clean_line,
                )
                doc.add_paragraph(item, style="List Number")
            else:
                doc.add_paragraph(clean_line)

    # Disclaimer
    doc.add_heading("Report Disclaimer", level=1)
    doc.add_paragraph(
        "AI-generated comparisons may contain errors or omissions. "
        "Verify important names, dates, amounts, deadlines, and "
        "requirements against the original PDFs. This report does not "
        "provide legal advice or certify compliance."
    )

    # Return the report as downloadable Word file bytes.
    output = BytesIO()
    doc.save(output)
    output.seek(0)

    return output.getvalue()



# ---------- SIDEBAR ----------

with st.sidebar:
    st.markdown(
        '<div class="brand">✦ NEXUS AI</div>',
        unsafe_allow_html=True,
    )
    st.caption("DOCUMENT INTELLIGENCE")
    st.divider()

    st.markdown("### Workspace")
    st.markdown("📊 Document Comparison")
    st.markdown("📁 Upload Center")
    st.markdown("📑 Report Export")

    st.divider()
    st.markdown("### System")
    st.markdown("🟣 PDF Text Extraction")
    st.markdown("🔵 Gemini AI Analysis")
    st.markdown("🟢 Word Report Generator")

    st.divider()
    st.caption("Your documents are processed in this session.")


# ---------- DASHBOARD ----------

st.markdown(
    '<div class="eyebrow">AI DOCUMENT INTELLIGENCE / WORKSPACE</div>',
    unsafe_allow_html=True,
)

st.markdown(
    '<div class="hero-title">Understand every '
    '<span class="gradient-text">difference.</span></div>',
    unsafe_allow_html=True,
)

st.markdown(
    '<p class="subtext">Compare documents intelligently. '
    'Discover meaningful changes and create clear reports.</p>',
    unsafe_allow_html=True,
)


# ---------- UPLOAD DOCUMENTS ----------

left, right = st.columns(2, gap="large")

with left:
    st.markdown(
        '<div class="panel"><h3>📄 Document A</h3>'
        '<p style="color:#A5AAC6">Original or previous version</p></div>',
        unsafe_allow_html=True,
    )
    original_file = st.file_uploader(
        "Choose Document A PDF",
        type=["pdf"],
        key="original_pdf",
    )

with right:
    st.markdown(
        '<div class="panel"><h3>📘 Document B</h3>'
        '<p style="color:#A5AAC6">Revised or updated version</p></div>',
        unsafe_allow_html=True,
    )
    revised_file = st.file_uploader(
        "Choose Document B PDF",
        type=["pdf"],
        key="revised_pdf",
    )


# ---------- EXTRACT, COUNT AND PREVIEW ----------

original_text = None
revised_text = None

if original_file is not None:
    if original_file.size == 0:
        st.error("Document A is an empty file.")
    elif original_file.size > 10 * 1024 * 1024:
        st.error("Document A exceeds the 10 MB file limit.")
    else:
        original_text, pages, characters, error = extract_pdf_text(
            original_file
        )

        if error:
            st.error(f"Document A: {error}")
        else:
            st.success(f"{original_file.name} read successfully!")

            metric1, metric2 = st.columns(2)
            metric1.metric("Document A Pages", pages)
            metric2.metric("Document A Characters", f"{characters:,}")

            with st.expander("Preview Document A extracted text"):
                st.text_area(
                    "Document A preview",
                    value=original_text[:5000],
                    height=220,
                    disabled=True,
                    key="original_preview",
                )
                if characters > 5000:
                    st.caption(
                        "Preview shows the first 5,000 characters. "
                        "The full extracted text is retained."
                    )

if revised_file is not None:
    if revised_file.size == 0:
        st.error("Document B is an empty file.")
    elif revised_file.size > 10 * 1024 * 1024:
        st.error("Document B exceeds the 10 MB file limit.")
    else:
        revised_text, pages, characters, error = extract_pdf_text(
            revised_file
        )

        if error:
            st.error(f"Document B: {error}")
        else:
            st.success(f"{revised_file.name} read successfully!")

            metric1, metric2 = st.columns(2)
            metric1.metric("Document B Pages", pages)
            metric2.metric("Document B Characters", f"{characters:,}")

            with st.expander("Preview Document B extracted text"):
                st.text_area(
                    "Document B preview",
                    value=revised_text[:5000],
                    height=220,
                    disabled=True,
                    key="revised_preview",
                )
                if characters > 5000:
                    st.caption(
                        "Preview shows the first 5,000 characters. "
                        "The full extracted text is retained."
                    )


# ---------- COMPARE BUTTON ----------

st.divider()

if st.button(
    "✦ Analyze & Compare Documents",
    type="primary",
    use_container_width=True,
):
    if original_file is None or revised_file is None:
        st.warning("Please upload both PDF documents.")
    elif not original_text or not revised_text:
        st.warning(
            "Both PDFs must contain readable text before comparison."
        )
    elif (
        len(original_text) > MAX_CHARS
        or len(revised_text) > MAX_CHARS
    ):
        st.warning(
            "One or both documents exceed 60,000 characters. "
            "Please use shorter PDFs. Text will not be silently cut off."
        )
    else:
        with st.spinner("Comparing documents with Gemini AI..."):
            result, error = compare_documents(
                original_text,
                revised_text,
            )

        if error:
            st.error(error)
        else:
            st.session_state["comparison_result"] = result
            st.session_state["original_name"] = original_file.name
            st.session_state["revised_name"] = revised_file.name
            st.success("Comparison completed successfully!")


# ---------- SPLIT COMPARISON SECTIONS ----------


def split_comparison_sections(result):
    heading_map = {
        "DOCUMENT A": "Document A",
        "DOCUMENT B": "Document B",
        "SIMILARITIES": "Similarities",
        "DIFFERENCES": "Differences",
        "KEY DETAILS COMPARISON": "Key Details Comparison",
        "OVERALL CONCLUSION": "Overall Conclusion",
        "LIMITATIONS": "Limitations",
    }

    sections = {name: [] for name in heading_map.values()}
    current_section = None

    # Check longer headings before shorter headings.
    heading_names = sorted(heading_map, key=len, reverse=True)

    for line in result.splitlines():
        cleaned = line.strip()

        # Remove Markdown heading markers and bold formatting.
        cleaned = re.sub(r"^#{1,6}\s*", "", cleaned)
        cleaned = cleaned.strip().strip("*_").strip()

        # Remove heading numbering, such as "1. " or "2) ".
        cleaned = re.sub(
            r"^\d+\s*[.)\-:]\s*",
            "",
            cleaned,
        ).strip()

        normalized = re.sub(r"\s+", " ", cleaned).upper()

        matched_heading = None

        for heading in heading_names:
            if normalized == heading:
                matched_heading = heading
                break

            # Also recognize headings such as:
            # DOCUMENT A: Summary
            # DOCUMENT A — Summary
            # DOCUMENT A (Summary)
            if normalized.startswith(heading):
                remainder = normalized[len(heading):]

                if re.match(r"^\s*[:\-–—(]", remainder):
                    matched_heading = heading
                    break

        if matched_heading:
            current_section = heading_map[matched_heading]
            continue

        if current_section:
            sections[current_section].append(line)

    return {
        name: "\n".join(lines).strip()
        for name, lines in sections.items()
    }


# ---------- RESULTS AND REPORT ----------

if st.session_state.get("comparison_result"):
    result = st.session_state["comparison_result"]
    sections = split_comparison_sections(result)

    st.markdown("## 📊 Comparison Results")

    tab_names = [
    "Document A",
    "Document B",
    "Similarities",
    "Differences",
    "Key Details Comparison",
    "Overall Conclusion",
    "Limitations",
]

    tabs = st.tabs(tab_names)

    for tab, name in zip(tabs, tab_names):
        with tab:
            content = sections.get(name, "")

            if content:
                st.markdown(content)
            else:
                st.warning(
                    f"No content detected for {name}. "
                    "Check the AI response headings."
                )

    st.divider()

    tab_documents, tab_export = st.tabs([
        "📑 Extracted Text",
        "⬇️ Word Report",
    ])

    with tab_documents:
        first, second = st.columns(2)

        with first:
            st.markdown("### Document A")
            st.caption(st.session_state["original_name"])

            if original_text:
                st.text_area(
                    "Document A extracted text",
                    value=original_text[:12000],
                    height=300,
                    disabled=True,
                    key="result_original_text",
                )
            else:
                st.info("Upload Document A again to preview its text.")

        with second:
            st.markdown("### Document B")
            st.caption(st.session_state["revised_name"])

            if revised_text:
                st.text_area(
                    "Document B extracted text",
                    value=revised_text[:12000],
                    height=300,
                    disabled=True,
                    key="result_revised_text",
                )
            else:
                st.info("Upload Document B again to preview its text.")

    with tab_export:
        try:
            report_bytes = create_word_report(
                st.session_state["original_name"],
                st.session_state["revised_name"],
                result,
            )

            st.success("Your Word report is ready.")

            st.download_button(
                label="⬇️ Download Comparison Report (.docx)",
                data=report_bytes,
                file_name="Nexus_AI_Comparison_Report.docx",
                mime=(
                    "application/vnd.openxmlformats-officedocument."
                    "wordprocessingml.document"
                ),
                use_container_width=True,
            )

        except Exception as error:
            st.error(f"Word report generation failed: {error}")

# ---------- RESET ----------

st.divider()

if st.button(
    "Reset and Compare Again",
    key="reset_comparison_button",
):
    st.session_state["comparison_result"] = None
    st.session_state["original_name"] = ""
    st.session_state["revised_name"] = ""

    for key in [
        "original_pdf",
        "revised_pdf",
        "original_preview",
        "revised_preview",
        "result_original_text",
        "result_revised_text",
    ]:
        st.session_state.pop(key, None)

    st.rerun()

st.caption(
    "Nexus AI • Python • Streamlit • Gemini • Word Report Generator"
)
