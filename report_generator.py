
from docx import Document
from io import BytesIO
from datetime import datetime


def generate_comparison_report(
    document1_name,
    document2_name,
    comparison_result
):
    """
    Generate a downloadable Word report.
    """

    doc = Document()

    # Report title
    doc.add_heading(
        "AI Document Comparison Report",
        level=0
    )

    doc.add_paragraph(
        f"Generated on: {datetime.now().strftime('%d-%m-%Y %H:%M')}"
    )

    # Document information
    doc.add_heading("Documents Compared", level=1)
    doc.add_paragraph(f"Document 1: {document1_name}")
    doc.add_paragraph(f"Document 2: {document2_name}")

    # AI comparison results
    doc.add_heading("Comparison Analysis", level=1)

    for line in comparison_result.splitlines():
        line = line.strip()

        if not line:
            continue

        if line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith(("- ", "* ")):
            doc.add_paragraph(line[2:], style="List Bullet")
        elif len(line) > 2 and line[0].isdigit() and line[1:3] in (". ", ") "):
            doc.add_paragraph(line[3:], style="List Number")
        else:
            doc.add_paragraph(line)

    # Save report in memory
    report_buffer = BytesIO()
    doc.save(report_buffer)
    report_buffer.seek(0)

    return report_buffer.getvalue()
